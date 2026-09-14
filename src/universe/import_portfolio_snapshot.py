from argparse import ArgumentParser
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

from src.database import get_connection
from src.universe.analysis_universe import get_security_id


PORTFOLIO_NAME = "Historical Portfolio"

TWO_PLACES = Decimal("0.01")


def to_decimal(value):
    if value is None or value == "":
        return None

    return Decimal(str(value))


def money(value):
    value = to_decimal(value)

    if value is None:
        return None

    return value.quantize(
        TWO_PLACES,
        rounding=ROUND_HALF_UP,
    )


def normalize_header(value):
    if value is None:
        return ""

    return " ".join(
        str(value)
        .strip()
        .upper()
        .replace("/", " ")
        .replace("-", " ")
        .split()
    )


def find_column(headers, names):
    normalized_names = {
        normalize_header(name)
        for name in names
    }

    for column, header in headers.items():
        if normalize_header(header) in normalized_names:
            return column

    return None


def find_snapshot_month(sheet):
    # Most files have the month in B2.
    # May 2026 has it in B1.
    for row in range(1, 4):
        for column in range(1, 5):
            value = sheet.cell(
                row=row,
                column=column,
            ).value

            if isinstance(value, datetime):
                return value.date().replace(day=1)

            if isinstance(value, date):
                return value.replace(day=1)

    raise ValueError(
        "Could not find snapshot month "
        "in the spreadsheet."
    )


def find_header_row(sheet):
    for row in range(1, 15):
        first_value = normalize_header(
            sheet.cell(
                row=row,
                column=1,
            ).value
        )

        if first_value == "SYMBOL":
            return row

    raise ValueError(
        "Could not find the holdings header row."
    )


def get_columns(sheet, header_row):
    headers = {
        column: sheet.cell(
            row=header_row,
            column=column,
        ).value
        for column in range(
            1,
            sheet.max_column + 1,
        )
    }

    symbol_column = find_column(
        headers,
        {"SYMBOL"},
    )

    description_column = find_column(
        headers,
        {"DESCRIPTION"},
    )

    price_column = find_column(
        headers,
        {
            "PRICE",
            "PRICE SHARE",
        },
    )

    shares_column = find_column(
        headers,
        {"SHARES"},
    )

    value_column = find_column(
        headers,
        {"VALUE"},
    )

    cost_column = find_column(
        headers,
        {
            "COST",
            "COST BASIS",
        },
    )

    gain_loss_column = find_column(
        headers,
        {
            "GAIN LOSS",
        },
    )

    required = {
        "SYMBOL": symbol_column,
        "DESCRIPTION": description_column,
        "PRICE": price_column,
        "SHARES": shares_column,
        "VALUE": value_column,
        "COST": cost_column,
        "GAIN/LOSS": gain_loss_column,
    }

    missing = [
        name
        for name, column in required.items()
        if column is None
    ]

    if missing:
        raise ValueError(
            "Missing required spreadsheet columns: "
            + ", ".join(missing)
        )

    return required


def read_spreadsheet(filename):
    path = Path(filename)

    if not path.exists():
        raise FileNotFoundError(
            f"Spreadsheet not found: {path}"
        )

    workbook = load_workbook(
        path,
        data_only=True,
    )

    sheet = workbook.active

    snapshot_month = find_snapshot_month(
        sheet
    )

    header_row = find_header_row(
        sheet
    )

    columns = get_columns(
        sheet,
        header_row,
    )

    cash = None
    source_total = None
    source_cost_basis = None
    source_gain_loss = None

    holdings = []

    for row in range(
        header_row + 1,
        sheet.max_row + 1,
    ):
        symbol_value = sheet.cell(
            row=row,
            column=columns["SYMBOL"],
        ).value

        description_value = sheet.cell(
            row=row,
            column=columns["DESCRIPTION"],
        ).value

        market_value = money(
            sheet.cell(
                row=row,
                column=columns["VALUE"],
            ).value
        )

        if (
            str(description_value or "")
            .strip()
            .lower()
            == "cash"
        ):
            cash = market_value
            continue

        symbol = str(
            symbol_value or ""
        ).strip().upper()

        if not symbol:
            # The source portfolio total appears
            # below the holdings in the VALUE column.
            if (
                market_value is not None
                and source_total is None
            ):
                source_total = market_value

                cost_value = money(
                    sheet.cell(
                        row=row + 1,
                        column=columns["COST"],
                    ).value
                ) if row < sheet.max_row else None

                gain_value = money(
                    sheet.cell(
                        row=row + 1,
                        column=columns["GAIN/LOSS"],
                    ).value
                ) if row < sheet.max_row else None

                source_cost_basis = cost_value
                source_gain_loss = gain_value

            continue

        price = to_decimal(
            sheet.cell(
                row=row,
                column=columns["PRICE"],
            ).value
        )

        shares = to_decimal(
            sheet.cell(
                row=row,
                column=columns["SHARES"],
            ).value
        )

        cost_basis = money(
            sheet.cell(
                row=row,
                column=columns["COST"],
            ).value
        )

        gain_loss = money(
            sheet.cell(
                row=row,
                column=columns["GAIN/LOSS"],
            ).value
        )

        if (
            price is None
            or shares is None
            or market_value is None
        ):
            raise ValueError(
                f"Incomplete holding data for {symbol}"
            )

        holdings.append(
            {
                "ticker": symbol,
                "description": str(
                    description_value or ""
                ).strip(),
                "price": price,
                "shares": shares,
                "market_value": market_value,
                "cost_basis": cost_basis,
                "gain_loss": gain_loss,
            }
        )

    if cash is None:
        raise ValueError(
            "Could not find cash value."
        )

    if source_total is None:
        raise ValueError(
            "Could not find portfolio total."
        )

    return {
        "snapshot_month": snapshot_month,
        "cash": cash,
        "source_total": source_total,
        "source_cost_basis": source_cost_basis,
        "source_gain_loss": source_gain_loss,
        "holdings": holdings,
        # Store only a neutral filename.
        "source_filename": (
            f"{snapshot_month:%Y-%m}.xlsx"
        ),
    }


def validate_snapshot(snapshot):
    holdings = snapshot["holdings"]

    holdings_market_value = sum(
        (
            row["market_value"]
            for row in holdings
        ),
        Decimal("0.00"),
    )

    calculated_total = money(
        holdings_market_value
        + snapshot["cash"]
    )

    source_total = money(
        snapshot["source_total"]
    )

    if calculated_total != source_total:
        raise ValueError(
            "Portfolio value does not match "
            "the spreadsheet. "
            f"Calculated: ${calculated_total:,.2f}; "
            f"Source: ${source_total:,.2f}"
        )

    calculated_cost_basis = None

    if all(
        row["cost_basis"] is not None
        for row in holdings
    ):
        calculated_cost_basis = money(
            sum(
                (
                    row["cost_basis"]
                    for row in holdings
                ),
                Decimal("0.00"),
            )
        )

    source_cost_basis = snapshot[
        "source_cost_basis"
    ]

    if (
        source_cost_basis is not None
        and calculated_cost_basis is not None
        and calculated_cost_basis
        != source_cost_basis
    ):
        raise ValueError(
            "Cost basis does not match "
            "the spreadsheet. "
            f"Calculated: "
            f"${calculated_cost_basis:,.2f}; "
            f"Source: "
            f"${source_cost_basis:,.2f}"
        )

    return {
        "holdings_market_value":
            holdings_market_value,
        "calculated_total":
            calculated_total,
        "calculated_cost_basis":
            calculated_cost_basis,
    }


def resolve_securities(snapshot):
    resolved = []

    missing = []

    for holding in snapshot["holdings"]:
        ticker = holding["ticker"]

        security_id = get_security_id(
            ticker
        )

        if security_id is None:
            missing.append(ticker)
            continue

        resolved.append(
            {
                **holding,
                "security_id": security_id,
            }
        )

    if missing:
        raise ValueError(
            "Securities not found in database: "
            + ", ".join(sorted(missing))
        )

    return resolved


def save_snapshot(
    snapshot,
    resolved_holdings,
):
    snapshot_month = snapshot[
        "snapshot_month"
    ]

    universe_name = (
        f"portfolio_history_"
        f"{snapshot_month:%Y_%m}"
    )

    universe_description = (
        "Historical Portfolio investment club "
        f"portfolio snapshot for "
        f"{snapshot_month:%B %Y}."
    )

    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO analysis_universes (
                    name,
                    description
                )
                VALUES (%s, %s)
                ON CONFLICT (name)
                DO UPDATE SET
                    description =
                        EXCLUDED.description,
                    updated_at =
                        CURRENT_TIMESTAMP
                RETURNING id;
                """,
                (
                    universe_name,
                    universe_description,
                ),
            )

            universe_id = cursor.fetchone()[0]

            cursor.execute(
                """
                DELETE FROM
                    analysis_universe_members
                WHERE universe_id = %s;
                """,
                (
                    universe_id,
                ),
            )

            for holding in resolved_holdings:
                cursor.execute(
                    """
                    INSERT INTO
                        analysis_universe_members (
                            universe_id,
                            security_id
                        )
                    VALUES (%s, %s)
                    ON CONFLICT DO NOTHING;
                    """,
                    (
                        universe_id,
                        holding[
                            "security_id"
                        ],
                    ),
                )

            cursor.execute(
                """
                INSERT INTO portfolio_snapshots (
                    portfolio_name,
                    snapshot_month,
                    cash,
                    total_portfolio_value,
                    total_cost_basis,
                    source_filename
                )
                VALUES (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s
                )
                ON CONFLICT (
                    portfolio_name,
                    snapshot_month
                )
                DO UPDATE SET
                    cash =
                        EXCLUDED.cash,
                    total_portfolio_value =
                        EXCLUDED.total_portfolio_value,
                    total_cost_basis =
                        EXCLUDED.total_cost_basis,
                    source_filename =
                        EXCLUDED.source_filename,
                    updated_at =
                        CURRENT_TIMESTAMP
                RETURNING id;
                """,
                (
                    PORTFOLIO_NAME,
                    snapshot_month,
                    snapshot["cash"],
                    snapshot["source_total"],
                    snapshot[
                        "source_cost_basis"
                    ],
                    snapshot[
                        "source_filename"
                    ],
                ),
            )

            snapshot_id = cursor.fetchone()[0]

            cursor.execute(
                """
                DELETE FROM
                    portfolio_snapshot_holdings
                WHERE snapshot_id = %s;
                """,
                (
                    snapshot_id,
                ),
            )

            for holding in resolved_holdings:
                cursor.execute(
                    """
                    INSERT INTO
                        portfolio_snapshot_holdings (
                            snapshot_id,
                            security_id,
                            ticker,
                            description,
                            price,
                            shares,
                            market_value,
                            cost_basis,
                            gain_loss
                        )
                    VALUES (
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s
                    );
                    """,
                    (
                        snapshot_id,
                        holding["security_id"],
                        holding["ticker"],
                        holding["description"],
                        holding["price"],
                        holding["shares"],
                        holding["market_value"],
                        holding["cost_basis"],
                        holding["gain_loss"],
                    ),
                )

        conn.commit()

    return (
        snapshot_id,
        universe_name,
    )


def import_snapshot(filename):
    snapshot = read_spreadsheet(
        filename
    )

    validation = validate_snapshot(
        snapshot
    )

    resolved_holdings = resolve_securities(
        snapshot
    )

    snapshot_id, universe_name = save_snapshot(
        snapshot,
        resolved_holdings,
    )

    print()
    print("HISTORICAL PORTFOLIO SNAPSHOT IMPORT")
    print("=" * 60)

    print(
        f"Snapshot month: "
        f"{snapshot['snapshot_month']:%B %Y}"
    )

    print(
        f"Holdings: "
        f"{len(resolved_holdings)}"
    )

    print(
        f"Holdings market value: "
        f"${validation['holdings_market_value']:,.2f}"
    )

    print(
        f"Cash: "
        f"${snapshot['cash']:,.2f}"
    )

    print(
        f"Portfolio total: "
        f"${snapshot['source_total']:,.2f}"
    )

    if snapshot["source_cost_basis"] is not None:
        print(
            f"Cost basis: "
            f"${snapshot['source_cost_basis']:,.2f}"
        )
    else:
        print(
            "Cost basis: "
            "not reported in source summary"
        )

    if snapshot["source_gain_loss"] is not None:
        print(
            f"Reported gain/loss: "
            f"${snapshot['source_gain_loss']:,.2f}"
        )
    else:
        print(
            "Reported gain/loss: "
            "not reported in source summary"
        )

    print(
        f"Snapshot ID: {snapshot_id}"
    )

    print(
        f"Universe: {universe_name}"
    )

    print(
        f"Stored source filename: "
        f"{snapshot['source_filename']}"
    )

    print()
    print("Import completed successfully.")


def main():
    parser = ArgumentParser(
        description=(
            "Import a Historical Portfolio "
            "monthly spreadsheet snapshot."
        )
    )

    parser.add_argument(
        "filename",
        help="Path to the XLSX snapshot file.",
    )

    args = parser.parse_args()

    import_snapshot(
        args.filename
    )


if __name__ == "__main__":
    main()